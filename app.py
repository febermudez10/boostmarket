
import os
import re
import requests

from datetime import datetime, timezone
from urllib.parse import urljoin
from io import BytesIO
from concurrent.futures import ThreadPoolExecutor, as_completed

import psycopg
from psycopg.rows import dict_row

from bs4 import BeautifulSoup
from openpyxl import load_workbook

from fastapi import FastAPI, Request, UploadFile, File
from fastapi.staticfiles import StaticFiles
from fastapi.responses import HTMLResponse, JSONResponse, Response
from fastapi.templating import Jinja2Templates


# ============================================================
# BOOSTMARKET - SOUTH DADE TOYOTA
# ============================================================

DATABASE_URL = os.environ.get("DATABASE_URL")
AUTO_DEV_API_KEY = os.environ.get("AUTO_DEV_API_KEY")

TIMEOUT = 30

HEADERS = {
    "User-Agent": "South-Dade-Toyota-Inventory-Manager/1.0"
}

SOURCE = (
    "https://www.southdadetoyota.com/"
    "llm/inventory/?type=used"
)

# ============================================================
# PRECIO MARKETPLACE - FORMULA ORIGINAL
# ============================================================
#
# Precio Marketplace = Precio VINCUE + $1,000
#
# NO se agregan:
# - Dealer Fee de $899
# - Electronic Filing Fee de $595
#
# ============================================================

MARKETPLACE_MARGIN = 1000


def marketplace_price_from_vincue(price):
    if price is None or price <= 0:
        return 0

    return price + MARKETPLACE_MARGIN


# ============================================================
# FASTAPI
# ============================================================

app = FastAPI(
    title="BoostMarket - South Dade Toyota"
)

app.mount(
    "/static",
    StaticFiles(directory="static"),
    name="static"
)

templates = Jinja2Templates(
    directory="templates"
)


# ============================================================
# DATABASE
# ============================================================

def db():
    if not DATABASE_URL:
        raise RuntimeError(
            "DATABASE_URL no está configurado en Render."
        )

    return psycopg.connect(
        DATABASE_URL,
        row_factory=dict_row
    )


def init_db():
    with db() as connection:
        with connection.cursor() as cur:

            cur.execute("""
                CREATE TABLE IF NOT EXISTS vehicles (
                    vin TEXT PRIMARY KEY,
                    vehicle TEXT,
                    condition TEXT,
                    mileage TEXT,
                    your_price DOUBLE PRECISION,
                    marketplace_price DOUBLE PRECISION,
                    url TEXT,
                    photo TEXT,
                    stock TEXT,
                    pic_count INTEGER DEFAULT 0,
                    status TEXT DEFAULT 'pending',
                    lead_count INTEGER DEFAULT 0,
                    last_seen TEXT,
                    updated_at TEXT
                )
            """)

            cur.execute("""
                ALTER TABLE vehicles
                ADD COLUMN IF NOT EXISTS stock TEXT
            """)

            cur.execute("""
                ALTER TABLE vehicles
                ADD COLUMN IF NOT EXISTS pic_count
                INTEGER DEFAULT 0
            """)

            cur.execute("""
                CREATE TABLE IF NOT EXISTS vehicle_photos (
                    vin TEXT NOT NULL,
                    position INTEGER NOT NULL,
                    url TEXT NOT NULL,
                    updated_at TEXT NOT NULL,
                    PRIMARY KEY (vin, position)
                )
            """)

            cur.execute("""
                CREATE INDEX IF NOT EXISTS
                idx_vehicle_photos_vin
                ON vehicle_photos(vin)
            """)

            cur.execute("""
                CREATE TABLE IF NOT EXISTS app_migrations (
                    migration_key TEXT PRIMARY KEY,
                    applied_at TEXT NOT NULL
                )
            """)

            # No reiniciar ni borrar los leads existentes.
            # Las migraciones anteriores permanecen guardadas.

        connection.commit()


# ============================================================
# UTILIDADES
# ============================================================

def money(value):
    if value is None:
        return None

    match = re.search(
        r"\$?\s*([\d,]+(?:\.\d{2})?)",
        str(value)
    )

    if not match:
        return None

    try:
        return float(
            match.group(1).replace(",", "")
        )
    except ValueError:
        return None


def normalize_vincue_photo_url(value):
    if not value:
        return ""

    url = str(value).strip()

    if not url.startswith(("http://", "https://")):
        return ""

    url = url.replace(
        "https://cdn-img.vincue.net/image/opt-sz160x/",
        "https://cdn-img.vincue.net/image/",
        1
    )

    url = url.replace(
        "http://cdn-img.vincue.net/image/opt-sz160x/",
        "http://cdn-img.vincue.net/image/",
        1
    )

    return url


def excel_cell_url(cell):
    if cell is None:
        return ""

    hyperlink = getattr(
        cell,
        "hyperlink",
        None
    )

    if hyperlink and hyperlink.target:
        url = normalize_vincue_photo_url(
            hyperlink.target
        )

        if url:
            return url

    return normalize_vincue_photo_url(
        cell.value
    )


# ============================================================
# WEBSITE - CONSULTA DE FOTO
# ============================================================

def get_primary_photo(detail_url):
    if not detail_url or detail_url == SOURCE:
        return ""

    try:
        response = requests.get(
            detail_url,
            headers=HEADERS,
            timeout=TIMEOUT
        )

        response.raise_for_status()

        soup = BeautifulSoup(
            response.text,
            "html.parser"
        )

        for attrs in (
            {"property": "og:image"},
            {"property": "og:image:secure_url"},
            {"name": "twitter:image"}
        ):
            meta = soup.find(
                "meta",
                attrs=attrs
            )

            if meta and meta.get("content"):
                url = urljoin(
                    response.url,
                    meta["content"].strip()
                )

                if url.startswith(("http://", "https://")):
                    return url

        image_src = soup.find(
            "link",
            rel="image_src"
        )

        if image_src and image_src.get("href"):
            url = urljoin(
                response.url,
                image_src["href"]
            )

            if url.startswith(("http://", "https://")):
                return url

        excluded = (
            "logo",
            "icon",
            "carfax",
            "pixel",
            "spinner",
            "placeholder",
            "avatar",
            "badge"
        )

        for image in soup.find_all("img"):
            raw = (
                image.get("data-src")
                or image.get("data-lazy-src")
                or image.get("data-original")
                or image.get("src")
                or ""
            ).strip()

            if not raw or raw.startswith("data:"):
                continue

            url = urljoin(
                response.url,
                raw
            )

            lower = url.lower()

            if any(word in lower for word in excluded):
                continue

            if any(
                extension in lower
                for extension in (
                    ".jpg",
                    ".jpeg",
                    ".png",
                    ".webp"
                )
            ):
                return url

    except requests.RequestException:
        pass

    return ""


# ============================================================
# WEBSITE COLLECTOR - LEGACY
# ============================================================
#
# Conservado por compatibilidad.
# No se ejecuta desde /api/sync.
# El inventario oficial sigue siendo VINCUE.
#
# ============================================================

def collect():
    response = requests.get(
        SOURCE,
        headers=HEADERS,
        timeout=TIMEOUT
    )

    response.raise_for_status()

    soup = BeautifulSoup(
        response.text,
        "html.parser"
    )

    found = {}

    for block in soup.find_all(
        ["article", "li", "div"]
    ):
        text = " ".join(
            block.stripped_strings
        )

        match = re.search(
            r"VIN:\s*([A-HJ-NPR-Z0-9]{17})",
            text,
            re.I
        )

        if not match:
            continue

        vin = match.group(1).upper()

        if vin in found:
            continue

        price_match = re.search(
            r"\$[\d,]+(?:\.\d{2})?",
            text
        )

        mileage_match = re.search(
            r"([\d,]+)\s*miles",
            text,
            re.I
        )

        vehicle_name = "Used vehicle"

        for tag in block.find_all(
            ["h2", "h3", "h4", "a"]
        ):
            title = " ".join(
                tag.stripped_strings
            )

            if (
                re.search(r"\b20\d{2}\b", title)
                and len(title) < 160
            ):
                vehicle_name = title
                break

        detail_url = SOURCE

        for link in block.find_all(
            "a",
            href=True
        ):
            url = urljoin(
                response.url,
                link["href"]
            )

            if (
                "southdadetoyota.com" in url
                and (
                    "inventory" in url
                    or "/used-" in url
                )
            ):
                detail_url = url
                break

        price = (
            money(price_match.group(0))
            if price_match
            else None
        )

        found[vin] = {
            "vin": vin,
            "vehicle": vehicle_name,
            "condition": "Used",
            "mileage": (
                mileage_match.group(1).replace(",", "")
                if mileage_match
                else ""
            ),
            "your_price": price,
            "marketplace_price": (
                marketplace_price_from_vincue(price)
                if price is not None
                else 0
            ),
            "url": detail_url,
            "photo": get_primary_photo(detail_url)
        }

    now = datetime.now(
        timezone.utc
    ).isoformat()

    with db() as connection:
        with connection.cursor() as cur:

            for vehicle in found.values():
                cur.execute("""
                    SELECT status, lead_count, photo
                    FROM vehicles
                    WHERE vin = %s
                """, (vehicle["vin"],))

                old = cur.fetchone()

                status = (
                    old["status"]
                    if old
                    else "pending"
                )

                leads = (
                    old["lead_count"]
                    if old
                    else 0
                )

                if (
                    not vehicle["photo"]
                    and old
                    and old.get("photo")
                ):
                    vehicle["photo"] = old["photo"]

                cur.execute("""
                    INSERT INTO vehicles (
                        vin,
                        vehicle,
                        condition,
                        mileage,
                        your_price,
                        marketplace_price,
                        url,
                        photo,
                        status,
                        lead_count,
                        last_seen,
                        updated_at
                    )
                    VALUES (
                        %s,%s,%s,%s,%s,%s,
                        %s,%s,%s,%s,%s,%s
                    )
                    ON CONFLICT (vin)
                    DO UPDATE SET
                        vehicle = EXCLUDED.vehicle,
                        condition = EXCLUDED.condition,
                        mileage = EXCLUDED.mileage,
                        your_price = EXCLUDED.your_price,
                        marketplace_price = EXCLUDED.marketplace_price,
                        url = EXCLUDED.url,
                        photo = EXCLUDED.photo,
                        last_seen = EXCLUDED.last_seen,
                        updated_at = EXCLUDED.updated_at
                """, (
                    vehicle["vin"],
                    vehicle["vehicle"],
                    vehicle["condition"],
                    vehicle["mileage"],
                    vehicle["your_price"],
                    vehicle["marketplace_price"],
                    vehicle["url"],
                    vehicle["photo"],
                    status,
                    leads,
                    now,
                    now
                ))

            if found:
                cur.execute("""
                    UPDATE vehicles
                    SET
                        status = 'unavailable',
                        updated_at = %s
                    WHERE NOT (vin = ANY(%s))
                """, (
                    now,
                    list(found.keys())
                ))

        connection.commit()

    return len(found)


# ============================================================
# STARTUP
# ============================================================

@app.on_event("startup")
def startup():
    init_db()


# ============================================================
# DASHBOARD
# ============================================================

@app.get(
    "/",
    response_class=HTMLResponse
)
def home(request: Request):
    return templates.TemplateResponse(
        request=request,
        name="dashboard.html"
    )


# ============================================================
# INVENTORY API
# ============================================================

@app.get("/api/inventory")
def inventory():
    with db() as connection:
        with connection.cursor() as cur:
            cur.execute("""
                SELECT *
                FROM vehicles
                ORDER BY updated_at DESC
            """)

            return cur.fetchall()


# ============================================================
# SUMMARY
# ============================================================

@app.get("/api/summary")
def summary():
    with db() as connection:
        with connection.cursor() as cur:

            cur.execute("""
                SELECT COUNT(*) AS n
                FROM vehicles
                WHERE status != 'unavailable'
            """)

            total = cur.fetchone()["n"]

            cur.execute("""
                SELECT COUNT(*) AS n
                FROM vehicles
                WHERE status = 'published'
            """)

            published = cur.fetchone()["n"]

            cur.execute("""
                SELECT COUNT(*) AS n
                FROM vehicles
                WHERE status = 'pending'
            """)

            pending = cur.fetchone()["n"]

            cur.execute("""
                SELECT COUNT(*) AS n
                FROM vehicles
                WHERE status = 'error'
            """)

            errors = cur.fetchone()["n"]

            cur.execute("""
                SELECT COALESCE(
                    SUM(lead_count),
                    0
                ) AS n
                FROM vehicles
            """)

            leads = cur.fetchone()["n"]

    return {
        "vehicles": total,
        "available": total,
        "published": published,
        "pending": pending,
        "errors": errors,
        "total_leads": leads
    }


# ============================================================
# WEBSITE SYNC - DESACTIVADO
# ============================================================

@app.post("/api/sync")
def sync():
    """
    No modifica el inventario.
    VINCUE sigue siendo la fuente de inventario.
    """
    try:
        with db() as connection:
            with connection.cursor() as cur:
                cur.execute("""
                    SELECT COUNT(*) AS n
                    FROM vehicles
                    WHERE status != 'unavailable'
                """)

                count = cur.fetchone()["n"]

        return {
            "ok": True,
            "count": count,
            "mode": "vincue",
            "message": (
                "La sincronización directa del website "
                "está desactivada. Usa Importar Vincue."
            )
        }

    except Exception as exc:
        return JSONResponse(
            {
                "ok": False,
                "error": str(exc)
            },
            status_code=500
        )


# ============================================================
# AUTO.DEV - OBTENER FOTOS
# ============================================================

def fetch_auto_dev_photos(vin):
    if not AUTO_DEV_API_KEY:
        return []

    try:
        response = requests.get(
            f"https://api.auto.dev/photos/{vin}",
            headers={
                "Authorization": (
                    f"Bearer {AUTO_DEV_API_KEY}"
                ),
                "Accept": "application/json"
            },
            timeout=15
        )

        response.raise_for_status()

        retail = (
            (response.json().get("data") or {})
            .get("retail")
            or []
        )

        photos = []
        seen = set()

        for url in retail:
            if (
                isinstance(url, str)
                and url.startswith(("http://", "https://"))
                and url not in seen
            ):
                seen.add(url)
                photos.append(url)

        return photos

    except Exception:
        return []


def fetch_photos_for_vins(
    vins,
    max_workers=8
):
    results = {}

    vins = list(
        dict.fromkeys(vins)
    )

    if not AUTO_DEV_API_KEY:
        return results

    with ThreadPoolExecutor(
        max_workers=max_workers
    ) as executor:

        jobs = {
            executor.submit(
                fetch_auto_dev_photos,
                vin
            ): vin
            for vin in vins
        }

        for job in as_completed(jobs):
            vin = jobs[job]

            try:
                results[vin] = job.result()
            except Exception:
                results[vin] = []

    return results


def save_vehicle_photos(
    cur,
    vin,
    photos,
    now
):
    if not photos:
        return 0

    cur.execute("""
        DELETE FROM vehicle_photos
        WHERE vin = %s
    """, (vin,))

    for position, url in enumerate(
        photos,
        1
    ):
        cur.execute("""
            INSERT INTO vehicle_photos (
                vin,
                position,
                url,
                updated_at
            )
            VALUES (%s,%s,%s,%s)
            ON CONFLICT (vin, position)
            DO UPDATE SET
                url = EXCLUDED.url,
                updated_at = EXCLUDED.updated_at
        """, (
            vin,
            position,
            url,
            now
        ))

    cur.execute("""
        UPDATE vehicles
        SET
            photo = %s,
            updated_at = %s
        WHERE vin = %s
    """, (
        photos[0],
        now,
        vin
    ))

    return len(photos)


# ============================================================
# IMPORTAR VINCUE
# ============================================================

@app.post("/api/import-vincue")
async def import_vincue(
    file: UploadFile = File(...)
):
    """
    Importa vehículos usados desde VINCUE.

    PRECIO MARKETPLACE:
        Precio VINCUE + $1,000

    Conserva leads, estados y fotografías.
    """
    try:
        contents = await file.read()

        workbook = load_workbook(
            BytesIO(contents),
            data_only=True
        )

        sheet = workbook.active

        headers = [
            str(cell.value).strip()
            if cell.value is not None
            else ""
            for cell in sheet[1]
        ]

        columns = {
            name: index
            for index, name in enumerate(headers)
        }

        required = [
            "VIN",
            "Year",
            "Model",
            "StockNo",
            "Odo",
            "Price"
        ]

        missing = [
            name
            for name in required
            if name not in columns
        ]

        if missing:
            return JSONResponse(
                {
                    "ok": False,
                    "error": (
                        "Faltan columnas: "
                        + ", ".join(missing)
                    )
                },
                status_code=400
            )

        photo_column = next(
            (
                name
                for name in [
                    "Photo",
                    "PhotoURL",
                    "PhotoUrl",
                    "Photo URL",
                    "Image",
                    "ImageURL",
                    "ImageUrl",
                    "Image URL",
                    "PrimaryPhoto",
                    "Primary Photo",
                    "PrimaryImage",
                    "Primary Image",
                    "Picture",
                    "PictureURL",
                    "Picture URL",
                    "Thumbnail",
                    "ThumbnailURL",
                    "Thumbnail URL"
                ]
                if name in columns
            ),
            None
        )

        imported = 0
        imported_vins = []

        now = datetime.now(
            timezone.utc
        ).isoformat()

        with db() as connection:
            with connection.cursor() as cur:

                for cells in sheet.iter_rows(
                    min_row=2,
                    values_only=False
                ):
                    row = [
                        cell.value
                        for cell in cells
                    ]

                    vin = str(
                        row[columns["VIN"]] or ""
                    ).strip().upper()

                    if not vin:
                        continue

                    year = str(
                        row[columns["Year"]] or ""
                    ).strip()

                    model = str(
                        row[columns["Model"]] or ""
                    ).strip()

                    stock = str(
                        row[columns["StockNo"]] or ""
                    ).strip()

                    odo = row[columns["Odo"]]

                    raw_price = row[
                        columns["Price"]
                    ]

                    try:
                        if isinstance(raw_price, str):
                            price = float(
                                raw_price.replace("$", "")
                                .replace(",", "")
                                .strip()
                                or 0
                            )
                        else:
                            price = float(
                                raw_price or 0
                            )
                    except (TypeError, ValueError):
                        price = 0

                    # ========================================
                    # PRECIO ORIGINAL DE BOOSTMARKET
                    # ========================================
                    #
                    # PRECIO VINCUE + $1,000
                    #
                    # No agregar Dealer Fee.
                    # No agregar Electronic Filing Fee.
                    #
                    # ========================================

                    marketplace_price = (
                        marketplace_price_from_vincue(
                            price
                        )
                    )

                    vehicle = (
                        f"{year} {model}"
                    ).strip()

                    pic_count = 0

                    if "PicCount" in columns:
                        try:
                            pic_count = int(
                                row[
                                    columns["PicCount"]
                                ] or 0
                            )
                        except (TypeError, ValueError):
                            pic_count = 0

                    photo = ""

                    if photo_column:
                        photo = excel_cell_url(
                            cells[
                                columns[photo_column]
                            ]
                        )

                    cur.execute("""
                        SELECT
                            status,
                            lead_count
                        FROM vehicles
                        WHERE vin = %s
                    """, (vin,))

                    old = cur.fetchone()

                    status = (
                        old["status"]
                        if old
                        and old["status"] != "unavailable"
                        else "pending"
                    )

                    leads = (
                        old["lead_count"]
                        if old
                        else 0
                    )

                    cur.execute("""
                        INSERT INTO vehicles (
                            vin,
                            vehicle,
                            condition,
                            mileage,
                            your_price,
                            marketplace_price,
                            url,
                            photo,
                            stock,
                            pic_count,
                            status,
                            lead_count,
                            last_seen,
                            updated_at
                        )
                        VALUES (
                            %s,%s,%s,%s,%s,%s,%s,
                            %s,%s,%s,%s,%s,%s,%s
                        )
                        ON CONFLICT (vin)
                        DO UPDATE SET
                            vehicle = EXCLUDED.vehicle,
                            condition = EXCLUDED.condition,
                            mileage = EXCLUDED.mileage,
                            your_price = EXCLUDED.your_price,
                            marketplace_price = EXCLUDED.marketplace_price,
                            stock = EXCLUDED.stock,
                            pic_count = EXCLUDED.pic_count,
                            status = EXCLUDED.status,
                            photo = CASE
                                WHEN EXCLUDED.photo IS NOT NULL
                                AND EXCLUDED.photo <> ''
                                THEN EXCLUDED.photo
                                ELSE vehicles.photo
                            END,
                            last_seen = EXCLUDED.last_seen,
                            updated_at = EXCLUDED.updated_at
                    """, (
                        vin,
                        vehicle,
                        "Used",
                        str(odo or ""),
                        price,
                        marketplace_price,
                        "",
                        photo,
                        stock,
                        pic_count,
                        status,
                        leads,
                        now,
                        now
                    ))

                    imported += 1
                    imported_vins.append(vin)

                # Actualizar disponibilidad solamente
                # cuando hay vehículos importados.
                if imported_vins:
                    cur.execute("""
                        UPDATE vehicles
                        SET
                            status = 'unavailable',
                            updated_at = %s
                        WHERE NOT (vin = ANY(%s))
                    """, (
                        now,
                        imported_vins
                    ))

            connection.commit()

        return {
            "ok": True,
            "count": imported,
            "source": "VINCUE",
            "marketplace_addition": MARKETPLACE_MARGIN,
            "message": (
                f"{imported} vehículos "
                "importados desde VINCUE"
            )
        }

    except Exception as exc:
        return JSONResponse(
            {
                "ok": False,
                "error": str(exc)
            },
            status_code=500
        )


# ============================================================
# PRUEBA DE PRECIO WEBSITE
# ============================================================

@app.get("/api/test-website-price/{vin}")
def test_website_price(vin: str):
    """
    Solo consulta.
    No modifica precios ni inventario.
    """
    vin = (vin or "").strip().upper()

    if not re.fullmatch(
        r"[A-HJ-NPR-Z0-9]{17}",
        vin
    ):
        return JSONResponse(
            {
                "ok": False,
                "error": "VIN inválido"
            },
            status_code=400
        )

    try:
        response = requests.get(
            SOURCE,
            headers=HEADERS,
            timeout=TIMEOUT
        )

        response.raise_for_status()

        soup = BeautifulSoup(
            response.text,
            "html.parser"
        )

        matching_blocks = []

        for tag in soup.find_all(
            ["article", "li", "div"]
        ):
            text = " ".join(
                tag.stripped_strings
            )

            if vin in text.upper():
                matching_blocks.append(tag)

        if not matching_blocks:
            return {
                "ok": False,
                "vin": vin,
                "website_status": (
                    response.status_code
                ),
                "message": (
                    "El website respondió, pero "
                    "no se encontró el VIN."
                )
            }

        block = min(
            matching_blocks,
            key=lambda tag: len(
                " ".join(tag.stripped_strings)
            )
        )

        text = " ".join(
            block.stripped_strings
        )

        prices = re.findall(
            r"\$\s*[\d,]+(?:\.\d{2})?",
            text
        )

        candidates = []

        for value in prices:
            amount = money(value)

            if amount is not None:
                candidates.append({
                    "display": value,
                    "amount": amount
                })

        return {
            "ok": True,
            "vin": vin,
            "website_status": (
                response.status_code
            ),
            "price_candidates": candidates,
            "website_text": text[:1800],
            "message": (
                "Consulta completada. "
                "Los precios encontrados "
                "requieren verificación."
            )
        }

    except requests.HTTPError as exc:
        status = (
            exc.response.status_code
            if exc.response is not None
            else 502
        )

        return {
            "ok": False,
            "vin": vin,
            "website_status": status,
            "message": (
                "El website rechazó la consulta. "
                "Se mantiene el cálculo "
                "VINCUE + $1,000."
            )
        }

    except requests.RequestException as exc:
        return {
            "ok": False,
            "vin": vin,
            "message": (
                "No se pudo consultar el website. "
                "Se mantiene VINCUE + $1,000."
            ),
            "detail": str(exc)
        }


# ============================================================
# PRUEBA FOTOS AUTO.DEV
# ============================================================

@app.get("/api/test-photos/{vin}")
def test_auto_dev_photos(vin: str):
    vin = (vin or "").strip().upper()

    if not re.fullmatch(
        r"[A-HJ-NPR-Z0-9]{17}",
        vin
    ):
        return JSONResponse(
            {
                "ok": False,
                "error": "VIN inválido"
            },
            status_code=400
        )

    if not AUTO_DEV_API_KEY:
        return JSONResponse(
            {
                "ok": False,
                "error": (
                    "AUTO_DEV_API_KEY "
                    "no está configurado"
                )
            },
            status_code=500
        )

    try:
        response = requests.get(
            f"https://api.auto.dev/photos/{vin}",
            headers={
                "Authorization": (
                    f"Bearer {AUTO_DEV_API_KEY}"
                ),
                "Accept": "application/json"
            },
            timeout=TIMEOUT
        )

        try:
            payload = response.json()
        except ValueError:
            payload = {
                "raw": response.text[:1000]
            }

        if response.status_code != 200:
            return JSONResponse(
                {
                    "ok": False,
                    "vin": vin,
                    "auto_dev_status": (
                        response.status_code
                    ),
                    "response": payload
                },
                status_code=response.status_code
            )

        data = (
            payload.get("data") or {}
            if isinstance(payload, dict)
            else {}
        )

        retail = (
            data.get("retail") or []
            if isinstance(data, dict)
            else []
        )

        photos = [
            url
            for url in retail
            if (
                isinstance(url, str)
                and url.startswith(
                    ("http://", "https://")
                )
            )
        ]

        return {
            "ok": True,
            "vin": vin,
            "count": len(photos),
            "photos": photos
        }

    except requests.RequestException as exc:
        return JSONResponse(
            {
                "ok": False,
                "vin": vin,
                "error": str(exc)
            },
            status_code=502
        )


# ============================================================
# SINCRONIZAR FOTOS AUTO.DEV
# ============================================================

@app.get("/api/sync-photos")
def sync_auto_dev_photos():
    if not AUTO_DEV_API_KEY:
        return JSONResponse(
            {
                "ok": False,
                "error": (
                    "AUTO_DEV_API_KEY "
                    "no está configurado"
                )
            },
            status_code=500
        )

    with db() as connection:
        with connection.cursor() as cur:
            cur.execute("""
                SELECT vin
                FROM vehicles
                WHERE status != 'unavailable'
                ORDER BY vin
            """)

            vins = [
                row["vin"]
                for row in cur.fetchall()
            ]

    results = fetch_photos_for_vins(
        vins
    )

    now = datetime.now(
        timezone.utc
    ).isoformat()

    total_photos = 0
    with_photos = 0
    missing = []

    with db() as connection:
        with connection.cursor() as cur:

            for vin in vins:
                photos = results.get(
                    vin,
                    []
                )

                if photos:
                    total_photos += save_vehicle_photos(
                        cur,
                        vin,
                        photos,
                        now
                    )

                    with_photos += 1
                else:
                    missing.append(vin)

        connection.commit()

    return {
        "ok": True,
        "vehicles_checked": len(vins),
        "vehicles_with_photos": with_photos,
        "total_photos": total_photos,
        "vehicles_without_photos": len(missing),
        "missing_vins": missing
    }


# ============================================================
# PUBLICACIONES
# ============================================================

@app.get(
    "/publicaciones",
    response_class=HTMLResponse
)
def publicaciones(request: Request):
    return templates.TemplateResponse(
        request=request,
        name="publicaciones.html"
    )


# ============================================================
# PHOTO PROXY - AUTO.DEV
# ============================================================

@app.get("/api/photo/{vin}/{position}")
def vehicle_photo(
    vin: str,
    position: int
):
    vin = (vin or "").strip().upper()

    if not re.fullmatch(
        r"[A-HJ-NPR-Z0-9]{17}",
        vin
    ):
        return Response(
            status_code=404
        )

    if position < 1:
        return Response(
            status_code=404
        )

    with db() as connection:
        with connection.cursor() as cur:
            cur.execute("""
                SELECT url
                FROM vehicle_photos
                WHERE vin = %s
                AND position = %s
            """, (
                vin,
                position
            ))

            row = cur.fetchone()

    if not row or not row.get("url"):
        return Response(
            status_code=404
        )

    try:
        headers = {
            "Accept": (
                "image/avif,image/webp,"
                "image/apng,image/svg+xml,"
                "image/*,*/*;q=0.8"
            ),
            "User-Agent": HEADERS["User-Agent"]
        }

        if AUTO_DEV_API_KEY:
            headers["Authorization"] = (
                f"Bearer {AUTO_DEV_API_KEY}"
            )

        response = requests.get(
            row["url"],
            headers=headers,
            timeout=TIMEOUT
        )

        response.raise_for_status()

        content_type = response.headers.get(
            "content-type",
            "image/jpeg"
        )

        if not content_type.lower().startswith(
            "image/"
        ):
            return Response(
                status_code=502
            )

        return Response(
            content=response.content,
            media_type=content_type.split(";")[0],
            headers={
                "Cache-Control": (
                    "public, max-age=86400"
                )
            }
        )

    except requests.RequestException:
        return Response(
            status_code=502
        )


# ============================================================
# VEHICLE DETAIL
# ============================================================

@app.get("/api/vehicle/{vin}")
def vehicle_detail(vin: str):
    vin = (vin or "").strip().upper()

    with db() as connection:
        with connection.cursor() as cur:
            cur.execute("""
                SELECT *
                FROM vehicles
                WHERE vin = %s
            """, (vin,))

            vehicle = cur.fetchone()

    if not vehicle:
        return JSONResponse(
            {
                "ok": False,
                "error": "Vehículo no encontrado"
            },
            status_code=404
        )

    with db() as connection:
        with connection.cursor() as cur:
            cur.execute("""
                SELECT position
                FROM vehicle_photos
                WHERE vin = %s
                ORDER BY position
            """, (vin,))

            photos = [
                (
                    f"/api/photo/{vin}/"
                    f"{item['position']}"
                )
                for item in cur.fetchall()
            ]

    return {
        "ok": True,
        "vehicle": vehicle,
        "photos": photos,
        "photo_count": len(photos)
    }


# ============================================================
# PREPARAR PUBLICACION
# ============================================================

@app.get(
    "/preparar/{vin}",
    response_class=HTMLResponse
)
def preparar_publicacion(
    request: Request,
    vin: str
):
    return templates.TemplateResponse(
        request=request,
        name="preparar.html"
    )


# ============================================================
# INVENTARIO
# ============================================================

@app.get(
    "/inventario",
    response_class=HTMLResponse
)
def inventario_page(request: Request):
    return templates.TemplateResponse(
        request=request,
        name="inventario.html"
    )


# ============================================================
# LEADS
# ============================================================

@app.get(
    "/leads",
    response_class=HTMLResponse
)
def leads_page(request: Request):
    return templates.TemplateResponse(
        request=request,
        name="leads.html"
    )


# ============================================================
# ACTIVIDAD
# ============================================================

@app.get(
    "/actividad",
    response_class=HTMLResponse
)
def actividad_page(request: Request):
    return templates.TemplateResponse(
        request=request,
        name="actividad.html"
    )


# ============================================================
# CONFIGURACION
# ============================================================

@app.get(
    "/configuracion",
    response_class=HTMLResponse
)
def configuracion_page(request: Request):
    return templates.TemplateResponse(
        request=request,
        name="configuracion.html"
    )
